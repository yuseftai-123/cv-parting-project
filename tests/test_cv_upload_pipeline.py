"""
Integration Tests for POST /api/v1/cvs/upload & 8-stage Pipeline (Phase 7 / Day 7)
"""
import os
import tempfile
import pytest
import docx
import fitz
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def _create_sample_docx(path: str):
    doc = docx.Document()
    doc.add_heading("CV - Youssef El Mansouri", level=1)
    doc.add_paragraph("Email: youssef.mansouri@gmail.com | Tél: +212 612345678 | CIN: AB123456")
    doc.add_paragraph("Né le 15/04/1992 à Casablanca | LinkedIn: linkedin.com/in/youssef-mansouri")
    
    doc.add_heading("Expérience Professionnelle", level=2)
    doc.add_paragraph("Développeur Full Stack Senior — Capgemini Technology Services")
    doc.add_paragraph("Janvier 2020 – Décembre 2023")
    doc.add_paragraph("Conception et développement d'APIs REST avec Python et FastAPI.")

    doc.add_heading("Formation", level=2)
    doc.add_paragraph("Master Génie Logiciel — Université Mohammed V, Rabat")
    doc.add_paragraph("2017 – 2019")
    doc.save(path)


def _create_sample_pdf(path: str):
    doc = fitz.open()
    page = doc.new_page()
    text = """
    Fatima Zahra Benali
    Analyste Business | Email: fatima.benali@company.ma | Tel: 0698765432
    CIN: CD987654 | Née le 20/09/1995 à Rabat
    
    Expérience Professionnelle
    Analyste Business — ENCG Consulting
    2021 – 2023
    Analyse des processus métier et rédaction de cahiers des charges.
    
    Formation
    Diplôme d'Ingénieur d'État — École Nationale des Sciences Appliquées (ENSA), Marrakech
    2014 – 2019
    """
    page.insert_text((50, 50), text)
    doc.save(path)
    doc.close()


def test_cv_upload_docx_pipeline_success():
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        tmp_path = tmp.name

    try:
        _create_sample_docx(tmp_path)
        with open(tmp_path, "rb") as f:
            response = client.post(
                "/api/v1/cvs/upload",
                files={"file": ("youssef_cv.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
            )
            
        assert response.status_code == 200, f"Upload failed: {response.text}"
        data = response.json()
        
        # Verify JSON Output Schema per Section 3.3
        assert "cv_id" in data
        assert "processing_timestamp" in data
        assert data["source_format"] == "docx"
        assert data["language_detected"] in ["fr", "en"]
        
        # Verify PII fields
        pii = data["pii_fields"]
        assert pii["email"]["masked"] is True
        assert pii["telephone"]["masked"] is True
        assert "youssef.mansouri@gmail.com" not in data["profil"]["resume_genere_llm"]
        
        # Verify experiences & formations
        assert len(data["experiences"]) >= 1
        assert len(data["formations"]) >= 1
        
        # Verify audit trail
        assert "pipeline_version" in data["audit_trail"]
        assert "layers_applied" in data["audit_trail"]
        assert data["audit_trail"]["processing_time_ms"] > 0

    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_cv_upload_pdf_pipeline_success():
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        tmp_path = tmp.name

    try:
        _create_sample_pdf(tmp_path)
        with open(tmp_path, "rb") as f:
            response = client.post(
                "/api/v1/cvs/upload",
                files={"file": ("fatima_cv.pdf", f, "application/pdf")}
            )
            
        assert response.status_code == 200, f"Upload failed: {response.text}"
        data = response.json()
        
        assert "cv_id" in data
        assert data["source_format"] == "pdf"
        assert data["language_detected"] == "fr"
        assert data["pii_fields"]["email"]["masked"] is True
        assert len(data["experiences"]) >= 1

    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_cv_upload_invalid_format_rejected():
    response = client.post(
        "/api/v1/cvs/upload",
        files={"file": ("invalid_file.txt", b"some plain text", "text/plain")}
    )
    assert response.status_code == 400
    assert "Unsupported file format" in response.json()["detail"]
