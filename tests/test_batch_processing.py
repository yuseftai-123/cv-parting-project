"""
Integration Tests for POST /api/v1/cvs/batch (Phase 8a)
"""
import os
import io
import zipfile
import tempfile
import pytest
import docx
import fitz
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def _create_sample_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_heading("CV - Tariq Mansouri", level=1)
    doc.add_paragraph("Email: tariq.mansouri@techcompany.ma | Tel: +212 655443322")
    doc.add_paragraph("CIN: CD987654 | Adresse: Casablanca")
    doc.add_heading("Expérience Professionnelle", level=2)
    doc.add_paragraph("DevOps Engineer — OCP Group (2020 - 2023)")
    doc.add_heading("Formation", level=2)
    doc.add_paragraph("Ingénieur d'État — ENSA (2011 - 2016)")
    
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _create_sample_pdf_bytes() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Fatima Benali\nEmail: fatima@domain.ma\nTel: 0612345678\nExpérience\nAnalyste — Deloitte\nFormation\nMaster - Univ Rabat")
    buffer = io.BytesIO()
    doc.save(buffer)
    doc.close()
    return buffer.getvalue()


def test_batch_upload_zip_success():
    """Test successful batch upload of a ZIP archive containing 2 valid CVs."""
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w") as zf:
        zf.writestr("cv1_tariq.docx", _create_sample_docx_bytes())
        zf.writestr("cv2_fatima.pdf", _create_sample_pdf_bytes())

    zip_bytes = zip_buffer.getvalue()

    response = client.post(
        "/api/v1/cvs/batch",
        files={"file": ("batch_cvs.zip", zip_bytes, "application/zip")}
    )

    assert response.status_code == 200, f"Batch upload failed: {response.text}"
    data = response.json()

    assert "batch_id" in data
    assert data["total_files"] == 2
    assert data["successful_files"] == 2
    assert data["failed_files"] == 0
    assert len(data["results"]) == 2
    assert data["total_processing_time_ms"] > 0


def test_batch_upload_mixed_valid_invalid():
    """Test batch upload containing 1 valid CV and 1 unsupported .txt file."""
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w") as zf:
        zf.writestr("valid_cv.docx", _create_sample_docx_bytes())
        zf.writestr("invalid_file.txt", b"plain text notes")

    zip_bytes = zip_buffer.getvalue()

    response = client.post(
        "/api/v1/cvs/batch",
        files={"file": ("mixed_batch.zip", zip_bytes, "application/zip")}
    )

    assert response.status_code == 200
    data = response.json()

    assert data["total_files"] == 2
    assert data["successful_files"] == 1
    assert data["failed_files"] == 1
    assert len(data["failed_details"]) == 1
    assert data["failed_details"][0]["filename"] == "invalid_file.txt"


def test_batch_upload_non_zip_rejected():
    """Test that uploading non-zip file returns 400 Bad Request."""
    response = client.post(
        "/api/v1/cvs/batch",
        files={"file": ("not_a_zip.pdf", b"%PDF-1.4 sample", "application/pdf")}
    )
    assert response.status_code == 400
    assert "Invalid batch archive format" in response.json()["detail"]
