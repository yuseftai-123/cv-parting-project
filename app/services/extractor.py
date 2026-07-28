import os
import fitz  # PyMuPDF
import docx

def extract_pdf(file_path: str) -> str:
    """
    Extracts text from a PDF file using PyMuPDF.
    If the extracted text is too short, detects it as a scanned/image PDF.
    """
    text_content = []
    
    try:
        # Open PDF file
        with fitz.open(file_path) as doc:
            for page in doc:
                text_content.append(page.get_text())
                
        full_text = "".join(text_content).strip()
        
        # Scanned PDF check: if we got almost no text, it is likely scanned
        if len(full_text) < 50:
            return "OCR not implemented"
            
        return full_text
    except Exception as e:
        print(f"Error reading PDF {file_path}: {e}")
        raise

def extract_docx(file_path: str) -> str:
    """
    Extracts text from a DOCX file using python-docx.
    Includes text inside tables to avoid missing structured content.
    """
    text_content = []
    
    try:
        doc = docx.Document(file_path)
        
        # Extract paragraph text
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_content.append(paragraph.text.strip())
                
        # Extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                # Deduplicate cell texts that might be repeated in merged cells
                clean_row_text = []
                for val in row_text:
                    if not clean_row_text or clean_row_text[-1] != val:
                        clean_row_text.append(val)
                if clean_row_text:
                    text_content.append(" | ".join(clean_row_text))
                    
        return "\n".join(text_content).strip()
    except Exception as e:
        print(f"Error reading DOCX {file_path}: {e}")
        raise

def extract_text(file_path: str) -> tuple[str, str]:
    """
    Unified extraction function.
    Given a file path, detects the format, extracts raw text,
    and returns a tuple of (raw_text, source_format).
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")
        
    _, ext = os.path.splitext(file_path.lower())
    
    if ext == ".pdf":
        text = extract_pdf(file_path)
        return text, "pdf"
    elif ext == ".docx":
        text = extract_docx(file_path)
        return text, "docx"
    else:
        raise ValueError(f"Unsupported file format: {ext}")
